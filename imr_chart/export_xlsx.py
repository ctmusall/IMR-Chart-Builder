"""Excel export of I-MR statistics, per-point data and out-of-limit rows."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .stats import ImrResult

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
FLAG_FILL = PatternFill("solid", fgColor="F8CBAD")


def stats_rows(result: ImrResult, column: str, source: str = "", excluded: int = 0) -> list[tuple[str, object]]:
    rows = [
        ("Source file", source),
        ("Column", column),
        ("Points (n)", result.n),
        ("Rows excluded", excluded),
        ("Mean (center line)", result.mean),
        ("Average moving range (MR̄)", result.mr_bar),
        ("Sigma (MR̄/1.128)", result.sigma),
        ("Sigma multiplier (k)", result.k),
        ("Individuals UCL", result.ucl),
        ("Individuals LCL", result.lcl),
        ("Moving range UCL", result.mr_ucl),
        ("Moving range LCL", result.mr_lcl),
        ("LSL", result.lsl if result.lsl is not None else "—"),
        ("Points above UCL", int(result.above_ucl.sum())),
        ("Points below LCL", int(result.below_lcl.sum())),
        ("Points below LSL", int(result.below_lsl.sum()) if result.lsl is not None else "—"),
        ("Moving ranges above MR UCL", int(result.mr_above_ucl.sum())),
    ]
    return rows


def points_frame(data: pd.DataFrame, result: ImrResult) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "Lot": data["label"].to_numpy() if "label" in data else range(1, result.n + 1),
            "Date": data["date"].dt.date.to_numpy() if "date" in data and data["date"].notna().all() else None,
            "Value": result.values,
            "Moving Range": result.moving_ranges,
            "Above UCL": result.above_ucl,
            "Below LCL": result.below_lcl,
            "Below LSL": result.below_lsl,
            "MR above UCL": result.mr_above_ucl,
        }
    )


def _style(ws, flag_cols: list[int] | None = None):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 45)
    ws.freeze_panes = "A2"
    if flag_cols:
        for row in ws.iter_rows(min_row=2):
            if any(row[c - 1].value is True for c in flag_cols):
                for cell in row:
                    cell.fill = FLAG_FILL


def save_xlsx(
    data: pd.DataFrame,
    result: ImrResult,
    column: str,
    path: str | Path,
    source: str = "",
    excluded: int = 0,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    points = points_frame(data, result)
    flagged = points[result.out_of_limits]
    stats = pd.DataFrame(stats_rows(result, column, source, excluded), columns=["Statistic", "Value"])
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        stats.to_excel(xw, sheet_name="Summary stats", index=False)
        points.to_excel(xw, sheet_name="Data", index=False)
        flagged.to_excel(xw, sheet_name="Out of limits", index=False)
        _style(xw.sheets["Summary stats"])
        _style(xw.sheets["Data"], flag_cols=[5, 6, 7, 8])
        _style(xw.sheets["Out of limits"])
    return path
