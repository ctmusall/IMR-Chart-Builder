import openpyxl
import pandas as pd

from imr_chart.export_xlsx import save_xlsx
from imr_chart.stats import compute_imr


def test_xlsx_export(tmp_path):
    vals = [10, 10, 10, 10, 10, 10, 10, 10, 10, 30]
    data = pd.DataFrame({"date": pd.date_range("2024-01-01", periods=10),
                         "label": [f"L{i}" for i in range(10)], "value": vals})
    r = compute_imr(vals, lsl=5)
    path = save_xlsx(data, r, "My Col", tmp_path / "x" / "out.xlsx", source="in.xlsx", excluded=4)

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Summary stats", "Verification", "Worked formulas", "Data", "Recompute",
                             "Input audit", "Out of limits"]
    stats = {row[0]: row[1] for row in wb["Summary stats"].iter_rows(min_row=2, values_only=True)}
    assert stats["Column"] == "My Col"
    assert stats["Source file"] == "in.xlsx"
    assert stats["Points (n)"] == 10
    assert stats["Rows excluded"] == 4
    assert stats["Mean (center line)"] == r.mean
    assert stats["Individuals UCL"] == r.ucl
    assert stats["LSL"] == 5
    assert stats["Points above UCL"] == 1
    assert stats["Points below LSL"] == 0

    data_ws = wb["Data"]
    assert data_ws.max_row == 11
    assert [c.value for c in data_ws[1]][:4] == ["Lot", "Date", "Value", "Moving Range"]
    assert data_ws["I1"].value == "Source row" and data_ws["I2"].value == 1
    oob = list(wb["Out of limits"].iter_rows(min_row=2, values_only=True))
    assert len(oob) == 1 and oob[0][0] == "L9"


def test_xlsx_export_without_lsl_or_dates(tmp_path):
    vals = [1.0, 2.0, 3.0]
    data = pd.DataFrame({"date": [pd.NaT] * 3, "label": ["a", "b", "c"], "value": vals})
    path = save_xlsx(data, compute_imr(vals), "C", tmp_path / "o.xlsx")
    stats = {row[0]: row[1] for row in openpyxl.load_workbook(path)["Summary stats"].iter_rows(min_row=2, values_only=True)}
    assert stats["LSL"] == "—"
    assert stats["Points below LSL"] == "—"


def test_xlsx_highlight_rule_column(tmp_path):
    vals = [10, 12, 11, 30]
    data = pd.DataFrame({"date": pd.date_range("2024-01-01", periods=4), "label": list("abcd"), "value": vals,
                         "rule": ["", "big", "", "big"], "highlight": ["", "#ff00ff", "", "#ff00ff"]})
    path = save_xlsx(data, compute_imr(vals), "V", tmp_path / "out.xlsx")
    ws = openpyxl.load_workbook(path)["Data"]
    header = [c.value for c in ws[1]]
    col = header.index("Highlight rule") + 1
    assert header[8:10] == ["Source row", "Source file"]  # Recompute's lookups are unaffected
    assert [ws.cell(r, col).value for r in range(2, 6)] == [None, "big", None, "big"]
    assert ws.cell(3, col).fill.fgColor.rgb.endswith("FF00FF")
