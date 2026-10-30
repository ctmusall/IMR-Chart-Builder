"""Several files with the same layout pooled into one chart."""

import shutil

import openpyxl
import pandas as pd
import pytest

from imr_chart import loader
from imr_chart.cli import main
from imr_chart.report import build_job, build_multi_job

from conftest import ROWS, _summary_workbook
from test_proof import checks, evaluate

COL, DATE = "QC3 CD34/Unit", "Date Processed"


def tables(paths):
    return [(p, loader.load_table(p)) for p in paths]


def test_pooled_matches_single_file(xlsx_split, xlsx_with_header_tab):
    pooled = build_multi_job(tables(xlsx_split), COL, DATE, lsl=2e7)
    single = build_job(loader.load_table(xlsx_with_header_tab), xlsx_with_header_tab, COL, DATE, lsl=2e7)
    assert pooled.result.n == single.result.n == 4
    assert pooled.excluded == single.excluded == 2
    assert pooled.result.ucl == pytest.approx(single.result.ucl)
    assert list(pooled.data["label"]) == ["L1", "L2", "L3", "L6"]  # sorted by date across files
    assert list(pooled.data["file"]) == ["part1.xlsx", "part1.xlsx", "part1.xlsx", "part2.xlsx"]
    assert list(pooled.data["row"]) == [2, 3, 1, 3]
    assert pooled.duplicates == 0


def test_row_order_keeps_file_order(xlsx_split):
    job = build_multi_job(tables(reversed(xlsx_split)), COL, None)
    assert list(job.data["label"]) == ["L6", "L3", "L1", "L2"]


def test_duplicates_across_files(xlsx_split, tmp_path):
    overlap = _summary_workbook(tmp_path / "overlap.xlsx", ROWS[2:4])  # repeats L2 (L4 has no value)
    job = build_multi_job(tables([*xlsx_split, overlap]), COL, DATE)
    assert job.result.n == 5
    assert job.duplicates == 2


def test_missing_column_names_file(xlsx_split, csv_file):
    other = csv_file.with_name("other.csv")
    other.write_text("QC Lot,Date Processed\nL1,2024-01-01\n")
    with pytest.raises(KeyError, match="other.csv"):
        build_multi_job(tables([xlsx_split[0], other]), COL, DATE)


def test_common_columns():
    assert loader.common_columns([["a", "b", "c"], ["c", "a"], ["a", "c", "d"]]) == ["a", "c"]
    assert loader.common_columns([]) == []


def test_same_name_in_two_folders(xlsx_split, tmp_path):
    (tmp_path / "copy").mkdir()
    twin = shutil.copy(xlsx_split[1], tmp_path / "copy" / "part1.xlsx")
    job = build_multi_job(tables([xlsx_split[0], twin]), COL, DATE)
    assert job.data["file"].nunique() == 2


@pytest.fixture
def pooled_workbook(xlsx_split, tmp_path):
    pytest.importorskip("formulas")
    job = build_multi_job(tables(xlsx_split), COL, DATE, lsl=2e7)
    return job, job.save_xlsx(tmp_path / "out" / "pooled.xlsx")


def test_pooled_workbook_passes_and_traces_files(pooled_workbook, xlsx_split):
    job, path = pooled_workbook
    values = evaluate(path)
    assert values[("VERIFICATION", "B1")].startswith("ALL ")
    assert set(checks(values).values()) == {"PASS"}

    wb = openpyxl.load_workbook(path)
    stats = dict(wb["Summary stats"].iter_rows(min_row=2, values_only=True))
    assert stats["Source file"] == "2 files" and stats["Source files"] == 2
    assert stats["Source 2 path"] == str(xlsx_split[1].resolve())
    assert {stats["Source 1 SHA-256"], stats["Source 2 SHA-256"]} == set(job.sha256s)
    data = list(wb["Data"].iter_rows(min_row=2, values_only=True))
    assert [(r[0], r[8], r[9]) for r in data][-1] == ("L6", 3, "part2.xlsx")


def test_pooled_lookup_uses_file_not_just_row(pooled_workbook, tmp_path):
    """Row 3 exists in both files; pointing a point at the other file's row 3 must fail."""
    _, path = pooled_workbook
    tampered = tmp_path / "tampered.xlsx"
    shutil.copy(path, tampered)
    wb = openpyxl.load_workbook(tampered)
    assert wb["Data"]["I5"].value == 3 and wb["Data"]["J5"].value == "part2.xlsx"
    wb["Data"]["J5"] = "part1.xlsx"
    wb.save(tampered)
    assert checks(evaluate(tampered))["Charted values that differ from their source row"] == "FAIL"


def test_html_hover_and_footnote(xlsx_split, tmp_path):
    job = build_multi_job(tables(xlsx_split), COL, DATE)
    html = job.save_html(tmp_path / "c.html").read_text(encoding="utf-8")
    assert "part2.xlsx" in html and "data from 2 files" in html
    texts = [t.get_text() for t in job.figure().texts]
    assert any("Data from 2 files" in t for t in texts)


def test_cli_multiple_files(xlsx_split, tmp_path, capsys):
    assert main([*map(str, xlsx_split), "--list-columns"]) == 0
    assert COL in capsys.readouterr().out
    assert main([*map(str, xlsx_split), "-c", COL, "--lsl", "2e7", "-o", str(tmp_path / "out")]) == 0
    assert "files=2 n=4 excluded=2" in capsys.readouterr().out
    assert len(list((tmp_path / "out").iterdir())) == 3


def test_cli_missing_sheet_names_file(xlsx_split, csv_file, tmp_path, capsys):
    other = tmp_path / "other.xlsx"
    pd.DataFrame(ROWS).to_excel(other, sheet_name="Elsewhere", header=False, index=False)
    assert main([str(xlsx_split[0]), str(other), "--sheet", "Summary", "-c", COL]) == 1
    assert "other.xlsx" in capsys.readouterr().err


def test_pooled_numeric_columns(xlsx_split):
    frames = [df for _, df in tables(xlsx_split)]
    assert COL not in loader.numeric_columns(frames[1])  # 1 number, 1 junk: under the threshold alone
    assert COL in loader.pooled_numeric_columns(frames)
    assert "Notes" not in loader.pooled_numeric_columns(frames)
    assert loader.pooled_numeric_columns([]) == []
