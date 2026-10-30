"""The workbook's evidence: Excel formulas must recompute the program's numbers, and catch tampering."""

import hashlib
import shutil

import openpyxl
import pandas as pd
import pytest

from imr_chart import loader, proof
from imr_chart.report import build_job
from imr_chart.stats import compute_imr

formulas = pytest.importorskip("formulas")


def evaluate(path):
    """Calculate a workbook like Excel would; returns {(sheet, cell): value}."""
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    out = {}
    for key, value in sol.items():
        sheet, _, cell = key.rpartition("!")
        if ":" not in cell:
            out[(sheet.rsplit("]", 1)[-1].strip("'").upper(), cell)] = value.value[0, 0]
    return out


def checks(values):
    rows = sorted(r for (s, c) in values if s == "VERIFICATION" and c.startswith("E") and c[1:].isdigit()
                  for r in [int(c[1:])] if r > 11)
    return {values[("VERIFICATION", f"A{r}")]: values[("VERIFICATION", f"E{r}")] for r in rows
            if ("VERIFICATION", f"A{r}") in values}


@pytest.fixture
def messy_csv(tmp_path):
    path = tmp_path / "messy.csv"
    path.write_text("QC Lot,Date Processed,CD34\n"
                    "L3,2024-04-18,7.0e7\nL1,2024-03-12,5.0e7\n\nL2,2024-04-01,9.0e7\nL4,2024-04-28,\n"
                    "L5,2024-05-08,bad\nL6,2024-05-18,1.0e7\nL7,,4.0e7\nL8,2024-05-20,50%\n"
                    'L9,2024-05-21,"1,000"\nL10,2024-05-22,2.6e8\n')
    return path


@pytest.fixture
def workbook(messy_csv, tmp_path):
    job = build_job(loader.load_table(messy_csv), messy_csv, "CD34", "Date Processed", lsl=2e7)
    return job, job.save_xlsx(tmp_path / "out" / "proof.xlsx")


def test_every_check_passes(workbook):
    _, path = workbook
    values = evaluate(path)
    assert values[("VERIFICATION", "B1")] == "ALL 18 CHECKS PASS"
    assert set(checks(values).values()) == {"PASS"}


def test_tampering_is_detected(workbook, tmp_path):
    _, path = workbook
    tampered = tmp_path / "tampered.xlsx"
    shutil.copy(path, tampered)
    wb = openpyxl.load_workbook(tampered)
    wb["Data"]["C3"] = wb["Data"]["C3"].value * 1.01  # one charted value no longer matches its source row
    wb.save(tampered)
    result = checks(evaluate(tampered))
    assert result["Charted values that differ from their source row"] == "FAIL"
    assert result["Mean (Individuals center line)"] == "FAIL"
    assert result["Rows read from source"] == "PASS"


def test_input_audit_and_provenance(workbook, messy_csv):
    job, path = workbook
    wb = openpyxl.load_workbook(path)
    audit = {r[1]: r for r in wb["Input audit"].iter_rows(min_row=2, values_only=True)}
    assert len(audit) == 10  # the blank line isn't a data row
    assert audit["L4"][6:] == ("No", "Value blank")
    assert audit["L5"][6:] == ("No", "Value not numeric")
    assert audit["L7"][6:] == ("No", "Date blank")
    assert audit["L2"][0] == 5  # CSV line number, counting the header and the blank line
    stats = dict(wb["Summary stats"].iter_rows(min_row=2, values_only=True))
    assert stats["Source SHA-256"] == hashlib.sha256(messy_csv.read_bytes()).hexdigest()
    assert stats["Date column"] == "Date Processed" and stats["Program version"].startswith("imr-chart ")
    worked = [r for r in wb["Worked formulas"].iter_rows(min_row=2, values_only=True)]
    assert ("Individuals UCL", "UCL = x̄ + k·σ̂") == worked[5][:2] and worked[5][3] == pytest.approx(job.result.ucl, rel=1e-15)


def test_source_rows_match_excel_rows(xlsx_with_header_tab):
    df = loader.load_table(xlsx_with_header_tab)
    assert list(df.index) == [1, 2, 3, 4, 5, 6]  # no header row on the data tab
    data, _ = loader.prepare_series(df, "QC3 CD34/Unit", "Date Processed", "QC Lot")
    assert dict(zip(data["label"], data["row"])) == {"L1": 2, "L2": 3, "L3": 1, "L6": 6}


def test_inline_header_rows_start_at_2(xlsx_inline_header):
    assert loader.load_table(xlsx_inline_header, "Data").index[0] == 2


def test_worked_steps_numbers():
    r = compute_imr([10.0, 12.0, 11.0, 15.0], lsl=9)
    steps = {q: (numbers, value) for q, _, numbers, value in proof.worked_steps(r)}
    assert steps["Mean, x̄ (Individuals center line)"] == ("48 / 4", 12.0)
    assert steps["Average moving range, MR̄ (MR center line)"] == ("7 / 3", pytest.approx(7 / 3))
    assert steps["Check: MR UCL at k = 3"][1] == pytest.approx(r.mr_ucl)
    assert "Points below LSL" in steps
    assert "Check: MR UCL at k = 3" not in {q for q, *_ in proof.worked_steps(compute_imr([1.0, 2.0], k=2))}


def test_file_sha256(tmp_path):
    f = tmp_path / "a.bin"
    f.write_bytes(b"abc")
    assert proof.file_sha256(f) == hashlib.sha256(b"abc").hexdigest()
    assert proof.file_sha256(tmp_path / "missing") is None


def test_audit_reasons_without_date():
    df = pd.DataFrame({"v": [1.0, None, "x", "  "]}, index=[2, 3, 4, 5])
    audit = loader.audit_rows(df, "v")
    assert list(audit["reason"]) == ["", "Value blank", "Value not numeric", "Value blank"]
    assert list(audit["row"]) == [2, 3, 4, 5]
